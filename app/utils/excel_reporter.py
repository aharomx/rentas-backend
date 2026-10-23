from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import func, and_
from datetime import datetime, date
from typing import Optional, List
import os 

from app.config import empresa
from app.models.movimiento_almacen import MovimientoAlmacen
from app.models.inventario_cliente import InventarioCliente
from app.models.solicitud_almacen import SolicitudAlmacen
from app.models.matenimiento_equipo import MantenimientoEquipo
from app.models.mantenimiento_refaccion import MantenimientoRefaccion
from app.models.articulo import Articulo


# =============== ESTILOS BASE ======================

COLOR_HEADER="1A5490"
COLOR_SUBHEADER="E8F0F8"
COLOR_TITULO="1A5490"

def _aplicar_estilo_titulo(ws, titulo:str, row:int=1):
    """ Aplica estilo al título del reporte """

    ws.merge_cells(f"A{row}:G(row)")
    cell=ws.cell(row=row, column=1,value=titulo)
    cell.font=Font(bold=True,size=14,color=COLOR_TITULO)
    cell.alignment=Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[row].height=25

def _aplicar_estilo_subtitulo(ws, empresa_nombre:str, row:int=2):
    """ Aplica estilo al subtitulo con datos de la empresa """

    ws.merge_cells(f"A{row}:G{row}")
    cell=ws.cell(row=row,column=1, value=empresa_nombre)
    cell.font=Font(bold=True, size=10)
    cell.alignment=Alignment(horizontal="center")

def _aplicar_estilo_fecha(ws, fecha_str:str, row:int=3):
    """ Aplica estilo a la fecha de generación """

    ws.merge_cells(f"A{row}:G{row}")
    cell=ws.cell(row=row, column=1, value=f"Generado: {fecha_str}")
    cell.font=Font(italic=True, size=9, color="666666")
    cell.alignment=Alignment(horizontal="center")

def _aplicar_estilo_header(ws, row:int,num_cols:int):
    """ Aplica estilo a la fila de encabezados """

    for col in range(1,num_cols+1):
        cell=ws.cell(row=row, column=col)
        cell.font=Font(bold=True, color="FFFFFF",size=10)
        cell.fill=PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        cell.border=_get_border()

    ws.row_dimensions[row].height=20

def _aplicar_estilo_body(ws, row_inicio:int, row_fin:int, num_cols:int):
    """ Aplica estilo a las filas de datos """

    for row in range(row_inicio, row_fin+1):
        for col in range(1,num_cols+1):
            cell=ws.cell(row=row, column=col)
            cell.font=Font(size=9)
            cell.border=_get_border()
            cell.alignment=Alignment(vertical="center",wrap_text=False)

        if (row-row_inicio)%2==1:
            for col in range(1,num_cols+1):
                ws.cell(row=row, column=col).fill=PatternFill(
                    start_color="F5F5F5",end_color="F5F5F5", fill_type="solid"
                )

def _get_border():
    """ Retorna un borde estándar. """

    thin=Side(border_style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def _ajustar_columnas(ws, num_cols:int):
    """ Ajusta el ancho de las columnas automáticamente """

    for col in range(1, num_cols+1):
        letra=get_column_letter(col)
        max_len=0
        for row in ws[letra]:
            if row.value:
                max_len=max(max_len,len(str(row.value)))
        ws.column_dimensions[letra].width=min(max_len+4,40)

def _encabezados_reporte(ws, subtitulo:str):
    """ Configura los 3 primeros renglones del reporte """

    _aplicar_estilo_titulo(ws,subtitulo,row=1)
    _aplicar_estilo_subtitulo(ws,empresa.EMPRESA_NOMBRE,row=2)
    _aplicar_estilo_fecha(ws,datetime.now().strftime("%d/%m/%Y %H:%M"),row=3)


# =================== REPORTE: TÓNERS POR CLIENTE ======================

def reporte_toners_por_cliente(
        db:Session,
        id_cliente:Optional[int]=None,
        fecha_inicio:Optional[date]=None,
        fecha_fin:Optional[date]=None,
        output_dir:str="reportes",
) -> str:
    """ Genera reporte de tóners en poder de cada cliente """

    os.makedirs(output_dir, exist_ok=True)

    query=db.query(InventarioCliente).filter(
        InventarioCliente.numero_serie_toner.isnot(None)
    )

    if id_cliente:
        query=query.filter(InventarioCliente.id_cliente==id_cliente)
    if fecha_inicio:
        query=query.filter(InventarioCliente.fecha_entrega >= fecha_inicio)
    if fecha_fin:
        query=query.filter(InventarioCliente.fecha_entrega <= fecha_fin)

    items=query.order_by(
        InventarioCliente.id_cliente, InventarioCliente.fecha_entrega.desc()
    ).all()

    wb= Workbook()
    ws= wb.active
    ws.title="Tóners por Cliente"

    _encabezados_reporte(ws, "REPORTE DE TÓNERS POR CLIENTE")

    # Encabezados
    headers = [
        "Cliente",
        "Código Artículo",
        "Artículo",
        "No. Serie",
        "Fecha Entrega",
        "Estado",
        "Observaciones"
    ]

    row_header=5
    for col, h in enumerate(headers, start=1):
        ws.cell(row=row_header, column=col, value=h)

    _aplicar_estilo_header(ws, row_header, len(headers))

    # Datos
    row=row_header+1
    for item in items:
        ws.cell(row=row,column=1,value=item.cliente_nombre or "-")
        ws.cell(row=row,column=2,value=item.articulo_codigo or "-")
        ws.cell(row=row,column=3,value=item.articulo_nombre or "-")
        ws.cell(row=row,column=4,value=item.numero_serie_toner or "-")
        ws.cell(row=row,column=5,value=item.fecha_entrega.strftime("%d/%m/%Y") if item.fecha_entrega else "-")
        ws.cell(row=row,column=6,value=item.estado)
        ws.cell(row=row,column=7,value=item.observaciones or "-")
        row += 1

    if row >= row_header +1:
        _aplicar_estilo_body(ws,row_header+1, row-1, len(headers))

    _ajustar_columnas(ws, len(headers))

    filename=f"{output_dir}/toners_por_cliente_{datetime.now().strftime('%Y%m%d_%H$M%S')}.xlsx"
    wb.save(filename)
    return filename

# ================= REPORTE: REFACCIONES EN MANTENIMIENTOS ===================

def reporte_refacciones_mantenimientos(
        db:Session,
        fecha_inicio:Optional[date]=None,
        fecha_fin:Optional[date]=None,
        id_equipo:Optional[int]=None,
        output_dir:str="reportes",
) ->str:
    """ Genera reporte de refacciones usadas en mantenimientos """

    os.makedirs(output_dir, exist_ok=True)

    from app.models.mantenimiento import Mantenimiento

    query=db.query(MantenimientoRefaccion).join(
        MantenimientoEquipo, MantenimientoRefaccion.id_mantenimiento_equipo==MantenimientoEquipo.id
    ).join(
        Articulo, MantenimientoRefaccion.id_articulo==Articulo.id
    )

    if fecha_inicio:
        query=query.filter(Mantenimiento.fecha_solicitud >= fecha_inicio)
    if fecha_fin:
        query=query.filter(Mantenimiento.fecha_solicitud <= fecha_fin)
    if id_equipo:
        query=query.filter(MantenimientoEquipo.id_equipo==id_equipo)

    items=query.order_by(Mantenimiento.fecha_solicitud.desc()).all()

    wb=Workbook()
    ws=wb.active
    ws.title="Refacciones Mantenimientos"

    _encabezados_reporte(ws, "REPORTE DE REFACCIONES EN MANTENIMIENTOS")

    headers = [
        "Folio Mant.",
        "Fecha",
        "Equipo (Serie)",
        "Artículo",
        "Cantidad",
        "No. Serie",
        "Observaciones"
    ]
    row_header=5

    for col, h in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=h)

    _aplicar_estilo_header(ws, row_header, len(headers))

    row=row_header+1

    for ref in items:
        mant_eq=ref.mantenimiento_equipo
        mant=mant_eq.mantenimiento if mant_eq else None
        eq=mant_eq.equipo if mant_eq else None

        ws.cell(row=row, column=1, value=mant.folio if mant else "-")
        ws.cell(row=row, column=2, value=mant.fecha_solicitud.strftime("%d/%m/%Y") if mant and mant.fecha_solicitud else "-")
        ws.cell(row=row, column=3, value=eq.numero_serie if eq else "-")
        ws.cell(row=row, column=4, value=ref.articulo.nombre if ref.articulo else "-")
        ws.cell(row=row, column=5, value=ref.cantidad)
        ws.cell(row=row, column=6, value=ref.numero_serie_toner or "-")
        ws.cell(row=row, column=7, value=ref.observaciones or  "-")
        row += 1

    if row > row_header + 1:
        _aplicar_estilo_body(ws, row_header+1, row-1, len(headers))

    _ajustar_columnas(ws, len(headers))

    filename=f"{output_dir}/refacciones_mantenimientos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return filename

# ===================== REPORTE MOVIMIENTOS DE ALMACÉN ========================

def reporte_movimientos_almacén(
        db:Session,
        fecha_inicio:Optional[date]=None,
        fecha_fin:Optional[date]=None,
        tipo:Optional[str]=None,
        id_articulo:Optional[int]=None,
        output_dir:str="reportes",
) -> str:
    """ Genera reporte de movimientos de almacén """

    os.makedirs(output_dir, exist_ok=True)

    query=db.query(MovimientoAlmacen).join(Articulo,MovimientoAlmacen,id_articulo==Articulo.id)

    if fecha_inicio:
        query=db.query.filter(MovimientoAlmacen.fecha_movimiento >= fecha_inicio)
    if fecha_fin:
        query=db.query.filter(MovimientoAlmacen.fecha_movimiento <= fecha_fin)
    if tipo:
        query=query.filter(MovimientoAlmacen.tipo==tipo)
    if id_articulo:
        query=query.filter(MovimientoAlmacen.id_articulo==id_articulo)

    items=query.order_by(MovimientoAlmacen.fecha_movimiento.desc()).all()

    wb=Workbook
    ws=wb.active
    ws.title="Movimientos"

    _encabezados_reporte(ws, "REPORTE DE MOVIMIENTOS DE ALMACÉN")

    headers=[
        "Fecha",
        "Tipo",
        "Código",
        "Artículo",
        "Cantidad",
        "Stock Antes",
        "Stock Después",
        "Referencia",
        "Usuario"
    ]

    row_header=5
    for col, h in enumerate(headers, start=1):
        ws.cell(row=row_header,start=1)

    _aplicar_estilo_header(ws,row_header,len(headers))

    row=row_header+1
    for mov in items:
        ws.cell(row=row, column=1, value=mov.fecha_movimiento.strftime("%d/%m/%Y %H:%M") if mov.fecha_movimiento else "-")
        ws.cell(row=row, column=2, value=mov.tipo.upper())
        ws.cell(row=row, column=3, value=mov.articulo_codigo or "-")
        ws.cell(row=row, column=4, value=mov.articulo_nombre or "-")
        ws.cell(row=row, column=5, value=mov.cantidad)
        ws.cell(row=row, column=6, value=mov.stock_anterior)
        ws.cell(row=row, column=7, value=mov.stock_nuevo)
        ws.cell(row=row, column=8, value=mov.referencia or "-")
        ws.cell(row=row, column=9, value=mov.usuario_nombre or "-")
        row += 1

    if row > row_header + 1:
        _aplicar_estilo_body(ws, row_header+1, row-1,len(headers))

    _ajustar_columnas(ws, len(headers))

    filename=f"{output_dir}/movimientos_almacen_{datetime.now().strftime('%HY%m%d_%H%M%S')}.xlsx"

    wb.save(filename)
    return filename


# =============== REPORTE STOCK CRÍTICO =======================

def reporte_stock_critico(
        db:Session,
        output_dir:str="reportes",
) -> str:
    """ Genera reporte de artículos con stock bajo o nulo """

    os.makedirs(output_dir, exist_ok=True)

    items = db.query(Articulo).filter(
        Articulo.activo==True,
        Articulo.stock_actual<=Articulo.stock_minimo,
    ).order_by(Articulo.stock_actual.desc()).all()

    wb=Workbook()
    ws=wb.active
    ws.title="Stock Crítico"

    _encabezados_reporte(ws,"REPORTE DE STOCK CRÍTICO")

    headers=[
        "Código",
        "Artículo",
        "Tipo",
        "Categoría",
        "Stock Actual",
        "Stock Mínimo",
        "Faltante"
    ]
    row_header=5
    for col, h in enumerate(headers, start=1):
        ws.cell(row=row_header, column=col, value=h)

    _aplicar_estilo_header(ws, row_header, len(headers))

    row=row_header+1
    for art in items:
        ws.cell(row=row, column=1, value=art.codigo)
        ws.cell(row=row, column=2, value=art.nombre)
        ws.cell(row=row, column=3, value=art.tipo or "-")
        ws.cell(row=row, column=4, value=art.categoria.nombre if art.categoria else "-")
        ws.cell(row=row, column=5, value=art.stock_actual)
        ws.cell(row=row, column=6, value=art.stock_minimo)
        ws.cell(row=row, column=7, value=art.stock_minimo-art.stock_actual)
        row += 1

    if row > row_header + 1:
        _aplicar_estilo_body(ws, row_header + 1, row - 1, len(headers))

        # Resaltar los que tienen stock 0
        for r in range(row_header+1, row):
            if ws.cell(row=r, column=5).value == 0:
                for c in range(1, len(headers)+1):
                    ws.cell(row=r, column=c).fill = PatternFill(
                        start_color="FFE0E0", end_color="FFE0E0", fill_type="solid"
                    )

    _ajustar_columnas(ws, len(headers))

    filename = f"{output_dir}/stock_critico_{datetime.now().strftime('%Y%M%d_%H%M%S')}.xlsx"
    wb.save(filename)
    return filename

# ===================== REPORTE: HISTORIAL DE REFACCIONES POR EQUIPO ======================

def reporte_historial_refacciones_equipos(
        db:Session,
        id_equipo:int,
        fecha_inicio:Optional[date]=None,
        fecha_fin:Optional[date]=None,
        output_dir:str="reportes",
) -> str:
    """ Genera rerpote de refacciones usadas en un equipo específico """

    os.makedirs(output_dir, exist_ok=True)

    from app.models.mantenimiento import Mantenimiento
    from app.models.equipo import Equipo

    equipo=db.query(Equipo).filter(Equipo.id == id_equipo).first()
    if not equipo:
        raise ValueError("Equipo no encontrado")

    query=db.query(MantenimientoRefaccion).join(
        MantenimientoEquipo, MantenimientoRefaccion.id_mantenimiento_equipo==Mantenimiento.id
    ).filter(MantenimientoEquipo.id_equipo==id_equipo)

    if fecha_inicio:
        query=query.filter(Mantenimiento.fecha_solicitud >= fecha_inicio)
    if fecha_fin:
        query=query.filter(Mantenimiento.fecha_solicitud <= fecha_fin)

    items=query.order_by(Mantenimiento.fecha_solicitud.desc().all())

    wb=Workbook
    ws=wb.active
    ws.title="Refacciones Equipo"

    _encabezados_reporte(ws, f"HISTORIAL DE REFACCIONES - EQUIPO {equipo.numero_serie}")

    ws.cell(row=4, column=1, value=f"Equipo: {equipo.numero_serie} - {equipo.modelo.nombre_modelo if equipo.modelo else '-'}")

    headers= [
        "Fecha",
        "Folio Mant.",
        "Tipo Mant.",
        "Articulo",
        "Cantidad",
        "No. Serie",
        "Observaciones"
    ]

    row_header=6
    for col,h in enumerate(headers, start=1):
        ws.cell(row=row_header, column=col, value=h)

    _aplicar_estilo_header(ws, row_header, len(headers))

    row = row_header + 1
    total_refacciones = 0
    for ref in items:
        mant_eq=ref.mantenimiento_equipo
        mant=mant_eq.mantenimiento if mant_eq else None

        ws.cell(row=row, column=1, value=mant.fecha_solicitud.strftime("%d/%m/%Y") if mant and mant.fecha_solicitud else "-")
        ws.cell(row=row, column=2, value=mant.folio if mant else "-")
        ws.cell(row=row, column=3, value=mant.tipo.upper() if mant else "-")
        ws.cell(row=row, column=4, value=ref.articulo.nombre if ref.articulo else "-")
        ws.cell(row=row, column=5, value=ref.cantidad)
        ws.cell(row=row, column=6, value=ref.numero_serie_toner or "-")
        ws.cell(row=row, column=7, value=ref.observaciones or "-")
        total_refacciones += ref.cantidad
        row += 1

    if row > row_header + 1:
        _aplicar_estilo_body(ws, row_header + 1, row - 1, len(headers))

    # Total al final
    ws.cell(row=row, column=4, value="TOTAL REFACCIONES:").font=Font(bold=True)
    ws.cell(row=row, column=5, value=total_refacciones).font = Font(bold=True)

    _ajustar_columnas(ws, len(headers))

    filename=f"{output_dir}/refacciones_equipo_{equipo.numero_serie}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    wb.save(filename)
    return filename



